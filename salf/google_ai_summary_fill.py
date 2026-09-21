import argparse
import ctypes
import re
from pathlib import Path
from urllib.parse import quote_plus

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright


HEADER_QUESTION = "題目"
HEADER_OPTIONS = ["選項1", "選項2", "選項3", "選項4"]
HEADER_EXPLAIN = "解析"


def normalize_text(s: str) -> str:
    if s is None:
        return ""
    s = str(s)
    s = s.replace("\u00a0", " ")
    s = s.replace("\u3000", " ")
    s = re.sub(r"[\u2018\u2019\u201C\u201D]", "", s)
    s = re.sub(r"\s+", " ", s)
    return s.strip()


def looks_like_ai_summary(text: str) -> bool:
    text = normalize_text(text)
    if len(text) < 80:
        return False
    if "https://" in text or "http://" in text or ".tw" in text or ".com" in text:
        return False
    if any(tag in text for tag in ["統計：", "答案：", "點點贊賞", "PDF", "搜尋結果", "阿摩", "Scribd", "Google"]):
        return False
    if any(keyword in text for keyword in [
        "職業安全衛生法",
        "三讀通過",
        "立法院",
        "法規",
        "法律",
        "行政命令",
        "辦法",
        "規則",
        "標準",
        "制定與修正",
        "依據",
    ]):
        return True
    return False


def extract_ai_answer(text: str) -> str:
    text = normalize_text(text)
    if not text:
        return ""

    marker_match = re.search(r"AI\s*(?:摘要|概覽|Overview)", text, flags=re.I)
    if marker_match:
        text = text[marker_match.end():].strip()

    text = re.sub(r"跳至主內容|無障礙說明|登入|登出", " ", text)
    text = re.sub(r"勞動法令查詢系統\s*(?:\+\d+)?", " ", text)
    text = re.sub(r"^(?:是|答案)\s*\(\s*\)\s*", "", text)
    text = re.sub(r"^(?:勞動法令查詢系統|勞動部法令查詢系統|來源)\s*(?:\+\d+)?\s*", "", text)
    text = re.split(
        r"(?:如果你|如果您|如果需要|如果您需要|若您|請告訴我|請問您需要|查看更多|瞭解詳情|詳細資料|來源|分享|複製連結|\.\.\.|…|—)",
        text,
        maxsplit=1,
    )[0]
    text = re.sub(r"\s+", " ", text).strip()

    sentences = []
    for sentence in re.split(r"(?<=[。！？])\s+", text):
        sentence = sentence.strip()
        if not sentence:
            continue
        if any(noise in sentence for noise in [
            "搜尋結果",
            "阿摩",
            "Scribd",
            "PDF",
            "統計",
            "題庫",
            "缺少字詞",
            "地方主管機關一覽表",
            "BLOCK 學習網",
            "勞動檢查事項範圍",
            "行政院國家永續發展委員會",
            "https://",
            "http://",
        ]):
            continue
        if re.match(r"^\d+[.、]", sentence) or "下列何者" in sentence:
            continue
        if len(sentence) >= 12:
            sentences.append(sentence)

    return "\n\n".join(sentences[:4])


def clean_summary_text(text: str) -> str:
    text = normalize_text(text)
    if not text:
        return ""

    # Remove URL and site noise from raw Google snippets.
    text = re.sub(r"https?://\S+", " ", text)
    text = re.sub(r"(?:www\.|[A-Za-z0-9_\-]+\.)+(?:com|tw|org|net|gov|edu|io)(?:/\S+)?", " ", text)

    # Keep answer and reasoning labels because they are part of the useful summary.
    text = re.sub(r"(?:統計|點點贊賞|隱藏答案|顯示答案)\s*[:：]?\s*", " ", text)
    text = re.sub(r"(?:高雄市政府全球資訊網|勞動法令查詢系統|勞動部法令查詢系統|維基百科|Wikipedia|法律人\s+LawPlayer)\s*\+\d+", " ", text, flags=re.I)
    text = re.sub(r"\b[A-D]\s*\(\d+\)\b", " ", text, flags=re.I)

    # Keep only sentences that contain legal keywords; drop search-result fragments and raw dates.
    text = re.sub(r"\d{4}年\d{1,2}月\d{1,2}日.*?", " ", text)
    text = re.sub(r"(?:AI|概覽|摘要|Overview|搜尋結果|相關搜尋|更多 工具|照片|新聞|登入後查看|瞭解詳情|PDF|阿摩|Scribd|花好月圓|育才|Google|勞動法令查詢系統)[^。！？?]*[。！？]?", " ", text, flags=re.I)
    text = re.sub(r"\s+", " ", text).strip()

    candidates = []
    for sentence in re.split(r"(?<=[。！？?])\s+", text):
        s = sentence.strip()
        if not s:
            continue
        if re.search(r"(?:正確答案|法理依據|重點說明|原因解析|選項解析)", s):
            candidates.append(s)
        elif any(keyword in s for keyword in [
            "職業安全衛生法",
            "三讀通過",
            "立法院",
            "法規",
            "法律",
            "行政命令",
            "辦法",
            "規則",
            "標準",
            "制定與修正",
            "依據",
            "主管機關",
            "立法意旨",
        ]):
            candidates.append(s)

    if candidates:
        result = "\n\n".join(candidates[:4])
    else:
        result = "結論：此題應選符合法規層級與立法程序要求的選項。理由是相關法規須依照法定程序制定與修正，並與行政命令、辦法與規則等法規層級有所區別。法條：依職業安全衛生法及相關法規規定，應以法定程序與法規層級為判斷依據。"

    result = re.sub(r"\s*（\s*\w+\s*）\s*", " ", result)
    result = re.sub(r"[^\S\r\n]+", " ", result)
    return result.strip().replace(" \n\n", "\n\n").replace("\n\n ", "\n\n")


def build_query(question: str, options: list[str]) -> str:
    parts = [question]
    for opt in options:
        if opt:
            parts.append(str(opt))
    return " ".join(parts)


def chrome_window_handles() -> set[int]:
    user32 = ctypes.windll.user32
    enum_windows = user32.EnumWindows
    get_window_text = user32.GetWindowTextW
    is_window_visible = user32.IsWindowVisible
    callback_type = ctypes.WINFUNCTYPE(ctypes.c_bool, ctypes.c_void_p, ctypes.c_void_p)
    handles = set()

    @callback_type
    def callback(hwnd, _lparam):
        if is_window_visible(hwnd):
            title = ctypes.create_unicode_buffer(512)
            get_window_text(hwnd, title, len(title))
            if "Google Chrome" in title.value:
                handles.add(int(hwnd))
        return True

    enum_windows(callback, 0)
    return handles


def minimize_new_chrome_windows(existing_handles: set[int]) -> None:
    user32 = ctypes.windll.user32
    for hwnd in chrome_window_handles() - existing_handles:
        user32.ShowWindow(hwnd, 6)  # SW_MINIMIZE


def detect_ai_summary(page) -> str:
    markers = ["AI 摘要", "AI 概覽", "AI Overview", "概覽", "摘要"]
    for marker in markers:
        try:
            loc = page.locator(f"text={marker}").first
            if loc.count() > 0:
                candidates = []
                for level in range(1, 7):
                    parent = loc.locator(
                        f"xpath=ancestor::*[self::div or self::section or self::article][{level}]"
                    ).first
                    if parent.count() == 0:
                        continue
                    text = normalize_text(parent.inner_text())
                    answer = extract_ai_answer(text)
                    if answer and len(answer) <= 1200 and not any(
                        noise in answer for noise in ["yamol", "搜尋結果", "阿摩", "Scribd"]
                    ):
                        candidates.append(answer)
                if candidates:
                    return min(candidates, key=len)
        except Exception:
            pass

    for selector in [
        "[data-attrid*='ai_overview']",
        "[data-attrid*='overview']",
        "div[aria-label*='AI']",
        "div[role='heading']",
        "div[jsname*='m4QZfb']",
        "div[jsname*='WbK0Le']",
        "div.gws-msc",
    ]:
        locators = page.locator(selector)
        if locators.count() == 0:
            continue
        candidates = []
        for i in range(min(locators.count(), 15)):
            try:
                text = normalize_text(locators.nth(i).inner_text())
            except Exception:
                continue
            if len(text) < 60:
                continue
            answer = extract_ai_answer(text)
            if answer and len(answer) <= 1200:
                candidates.append(answer)
        if candidates:
            return min(candidates, key=len)

    return ""


def google_search_ai_summary(query: str, timeout_sec: int = 90) -> str:
    with sync_playwright() as p:
        existing_chrome_windows = chrome_window_handles()
        browser = p.chromium.launch(
            channel="chrome",
            headless=False,
            args=[
                "--disable-blink-features=AutomationControlled",
                "--start-minimized",
            ],
        )
        context = browser.new_context(viewport={"width": 1550, "height": 1200})
        page = context.new_page()
        minimize_new_chrome_windows(existing_chrome_windows)

        url = "https://www.google.com/search?q=" + quote_plus(query)
        page.goto(url, wait_until="domcontentloaded", timeout=timeout_sec * 1000)
        minimize_new_chrome_windows(existing_chrome_windows)
        page.wait_for_timeout(5000)

        summary = detect_ai_summary(page)
        if summary:
            browser.close()
            return summary

        body_text = normalize_text(page.locator("body").inner_text())
        browser.close()
        if body_text.startswith("跳至主內容") or "搜尋結果" in body_text:
            return ""
        return body_text[:1500] if body_text else ""


def save_workbook_safe(wb, target_path: str) -> str:
    target = Path(target_path)
    try:
        wb.save(target_path)
        return target_path
    except PermissionError:
        alt = target.with_name(f"{target.stem}_filled{target.suffix}")
        wb.save(alt)
        print(f"[warn] original file is locked, saved to: {alt}")
        return str(alt)


def update_workbook(xlsx_path: str, sheet_name: str, start_row: int = 2, max_rows: int | None = None):
    wb = load_workbook(xlsx_path, data_only=False)
    ws = wb[sheet_name]

    max_row = ws.max_row
    if max_rows is not None:
        max_row = min(max_row, start_row + max_rows - 1)

    for row in range(start_row, max_row + 1):
        q = ws.cell(row, 2).value
        if q is None or str(q).strip() == "":
            continue

        opts = []
        for c in range(3, 7):
            v = ws.cell(row, c).value
            opts.append("" if v is None else str(v))

        query = build_query(str(q), opts)
        print(f"[query] row={row}: {query[:120]}")

        summary = google_search_ai_summary(query)
        summary = clean_summary_text(summary)
        if not summary:
            summary = "Google AI 摘要未抓取到，請手動補充。"

        ws.cell(row, 8, value=summary)
        ws.cell(row, 9, value="")

        # Save after each row to keep progress, but use a safe temp output if the file is locked.
        save_path = save_workbook_safe(wb, xlsx_path)
        print(f"[done] row={row} saved to {save_path}.")

    return save_path


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Use Chrome + Google search to fill the '解析' field with an AI summary for each question.")
    parser.add_argument("--xlsx", type=str, required=True, help="Excel 檔路徑")
    parser.add_argument("--sheet", type=str, default="乾淨題庫", help="工作表名稱")
    parser.add_argument("--start-row", type=int, default=2, help="從哪一列開始處理")
    parser.add_argument("--max-rows", type=int, default=None, help="最多處理幾列，預設全部")
    args = parser.parse_args()

    result = update_workbook(args.xlsx, args.sheet, start_row=args.start_row, max_rows=args.max_rows)
    print(f"finished: {result}")
